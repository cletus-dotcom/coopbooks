"""Excel template generation and cooperative data migration import."""

import io
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash

from app import db
from app.config import is_valid_user_role, normalize_role
from app.migration_config import (
    ACCOUNT_TYPES,
    IMPORT_ORDER,
    LEDGER_TYPES,
    MEMBER_STATUSES,
    MIGRATION_TABLE_MAP,
    MIGRATION_TABLES,
    NORMAL_BALANCES,
)
from app.models import (
    Account,
    Cooperative,
    JournalEntry,
    JournalLine,
    Member,
    MemberLedger,
    User,
)
from app.member_registry_config import (
    CIVIL_STATUSES as REGISTRY_CIVIL_STATUSES,
    EDUCATION_LEVELS as REGISTRY_EDUCATION_LEVELS,
    GENDERS as REGISTRY_GENDERS,
    MEMBERSHIP_TYPES as REGISTRY_MEMBERSHIP_TYPES,
    build_full_name,
)
from app.member_service import create_opening_ledger_entries
from app.tenant_manager import get_coop_registry


def _header_style():
    return {
        "font": Font(bold=True, color="FFFFFF"),
        "fill": PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid"),
    }


def build_template_workbook(table_key):
    """Build a single-table Excel import template."""
    spec = MIGRATION_TABLE_MAP[table_key]
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]

    style = _header_style()
    for col, header in enumerate(spec["headers"], start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = style["font"]
        cell.fill = style["fill"]

    if spec.get("example"):
        for col, value in enumerate(spec["example"], start=1):
            ws.cell(row=2, column=col, value=value)

    ws.freeze_panes = "A2"
    for col in range(1, len(spec["headers"]) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    notes = wb.create_sheet("_instructions")
    notes["A1"] = spec["label"]
    notes["A1"].font = Font(bold=True, size=12)
    notes["A3"] = "PostgreSQL table"
    notes["A4"] = spec.get("db_table", spec["key"])
    notes["A6"] = "Description"
    notes["A7"] = spec["description"]
    notes["A9"] = "Required columns"
    notes["A10"] = ", ".join(spec.get("required", []))
    notes["A12"] = "Paste your data starting row 2 on the main sheet. Do not rename header row."
    valid = spec.get("valid_values") or {}
    if valid:
        notes["A14"] = "Valid values"
        row = 15
        for field, choices in valid.items():
            notes[f"A{row}"] = f"{field}: {', '.join(choices)}"
            row += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_all_templates_zip():
    """Zip archive containing all migration templates."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for spec in IMPORT_ORDER:
            wb_buf = build_template_workbook(spec["key"])
            zf.writestr(spec["filename"], wb_buf.getvalue())
        readme = (
            "CoopBooks Migration Templates\n"
            "=============================\n\n"
            "Single cooperative database — import order:\n"
        )
        for spec in IMPORT_ORDER:
            readme += f"  {spec['order']}. {spec['filename']} — {spec['label']} ({spec.get('db_table', spec['key'])})\n"
        readme += (
            "\nFill each sheet from row 2. Keep column headers unchanged.\n"
            "Upload via Administration → Migration Tool (PlatformAdmin role).\n"
        )
        zf.writestr("README.txt", readme)
    buffer.seek(0)
    return buffer


def _parse_bool(value):
    if value is None or str(value).strip() == "":
        return True
    return str(value).strip().lower() in ("1", "true", "yes", "y")


def _parse_decimal(value, default=0):
    if value is None or str(value).strip() == "":
        return Decimal(str(default))
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        raise ValueError(f"Invalid number: {value}")


def _parse_date(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date: {value} (use YYYY-MM-DD)")


def _row_dict(headers, row_values):
    data = {}
    for i, header in enumerate(headers):
        if i < len(row_values):
            val = row_values[i]
            if val is not None and str(val).strip() != "":
                data[header] = val
    return data


def _parse_int(value, default=0):
    if value is None or str(value).strip() == "":
        return default
    try:
        return int(str(value).strip())
    except ValueError:
        raise ValueError(f"Invalid integer: {value}")


def _parse_member_names(row):
    """Resolve MC 2012-16 name parts; supports legacy full_name column."""
    last_name = str(row.get("last_name") or "").strip()
    first_name = str(row.get("first_name") or "").strip()
    middle_name = str(row.get("middle_name") or "").strip() or None

    if row.get("full_name") and (not last_name or not first_name):
        full = str(row["full_name"]).strip()
        if "," in full:
            last_name, rest = [part.strip() for part in full.split(",", 1)]
            name_parts = rest.split(None, 1)
            first_name = name_parts[0]
            if len(name_parts) > 1:
                middle_name = name_parts[1]
        elif not first_name:
            first_name = full

    if not last_name or not first_name:
        raise ValueError("last_name and first_name are required")

    return last_name, first_name, middle_name


def _validate_member_row(row, row_num=None):
    """Validate a member import row against registry choices."""
    prefix = f"Row {row_num}: " if row_num else ""
    errors = []

    status = str(row.get("status", "Active")).strip()
    if status not in MEMBER_STATUSES:
        errors.append(f"{prefix}status must be Active, Inactive, or Terminated")

    gender = str(row.get("gender", "")).strip()
    if gender and gender not in REGISTRY_GENDERS:
        errors.append(f"{prefix}invalid gender '{gender}'")

    civil = str(row.get("civil_status", "")).strip()
    if civil and civil not in REGISTRY_CIVIL_STATUSES:
        errors.append(f"{prefix}invalid civil_status '{civil}'")

    education = str(row.get("highest_education", "")).strip()
    if education and education not in REGISTRY_EDUCATION_LEVELS:
        errors.append(f"{prefix}invalid highest_education '{education}'")

    mtype = str(row.get("membership_type", "Regular")).strip()
    if mtype and mtype not in REGISTRY_MEMBERSHIP_TYPES:
        errors.append(f"{prefix}invalid membership_type '{mtype}'")

    if status == "Terminated":
        if not str(row.get("termination_date") or "").strip():
            errors.append(f"{prefix}termination_date required when status is Terminated")
        if not str(row.get("termination_bod_resolution") or "").strip():
            errors.append(f"{prefix}termination_bod_resolution required when status is Terminated")

    return errors


def read_upload_rows(file_storage, table_key):
    """Parse uploaded Excel into list of row dicts."""
    spec = MIGRATION_TABLE_MAP[table_key]
    wb = load_workbook(file_storage, read_only=True, data_only=True)
    sheet_name = spec["sheet"]
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["File is empty."]

    headers = [str(h).strip() if h else "" for h in rows[0]]
    missing = [h for h in spec["headers"] if h not in headers]
    if missing:
        return [], [f"Missing columns: {', '.join(missing)}"]

    col_index = {h: headers.index(h) for h in spec["headers"]}
    parsed = []
    errors = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        values = [row[col_index[h]] if col_index[h] < len(row) else None for h in spec["headers"]]
        data = _row_dict(spec["headers"], values)
        for req in spec.get("required", []):
            if req not in data:
                errors.append(f"Row {row_num}: missing required '{req}'")
                break
        else:
            data["_row"] = row_num
            parsed.append(data)

    return parsed, errors


def validate_import(table_key, rows):
    """Dry-run validation without writing to database."""
    if not get_coop_registry():
        return {"valid_rows": 0, "errors": ["No cooperative registered."], "warnings": []}

    errors = []
    warnings = []

    if table_key == "accounts":
        for row in rows:
            atype = str(row.get("account_type", "")).strip()
            if atype not in ACCOUNT_TYPES:
                errors.append(f"Row {row['_row']}: invalid account_type '{atype}'")
            nb = str(row.get("normal_balance", "")).strip()
            if nb not in NORMAL_BALANCES:
                errors.append(f"Row {row['_row']}: normal_balance must be Debit or Credit")

    if table_key == "users":
        for row in rows:
            role = normalize_role(row.get("role", "Member"))
            if not is_valid_user_role(role):
                errors.append(f"Row {row['_row']}: invalid role '{row.get('role')}'")

    if table_key == "members":
        for row in rows:
            errors.extend(_validate_member_row(row, row.get("_row")))

    if table_key == "journal_lines":
        entry_nos = {e.entry_no for e in JournalEntry.query.all()}
        codes = {a.code for a in Account.query.all()}
        for row in rows:
            if row.get("entry_no") not in entry_nos:
                warnings.append(
                    f"Row {row['_row']}: entry_no '{row.get('entry_no')}' not in database yet "
                    "(import journal_entries first)"
                )
            if row.get("account_code") not in codes:
                errors.append(f"Row {row['_row']}: unknown account_code '{row.get('account_code')}'")

    if table_key == "member_ledger":
        member_nos = {m.member_no for m in Member.query.all()}
        entry_nos = {e.entry_no for e in JournalEntry.query.all()}
        for row in rows:
            if row.get("member_no") not in member_nos:
                errors.append(f"Row {row['_row']}: unknown member_no '{row.get('member_no')}'")
            ltype = str(row.get("ledger_type", "")).strip().lower()
            if ltype not in LEDGER_TYPES:
                errors.append(f"Row {row['_row']}: ledger_type must be share, savings, or loan")
            if row.get("entry_no") and row.get("entry_no") not in entry_nos:
                warnings.append(
                    f"Row {row['_row']}: entry_no '{row.get('entry_no')}' not found "
                    "(import journal_entries first or leave blank)"
                )

    return {"valid_rows": len(rows), "errors": errors, "warnings": warnings}


def import_rows(table_key, rows, mode="append"):
    """Import parsed rows into the cooperative database. mode: append | replace."""
    if not get_coop_registry():
        return {"table": table_key, "imported": 0, "skipped": 0, "errors": ["No cooperative registered."]}
    spec = MIGRATION_TABLE_MAP[table_key]
    imported = 0
    skipped = 0
    errors = []

    if mode == "replace":
        _clear_table(table_key)

    try:
        for row in rows:
            try:
                created = _import_one_row(table_key, row, mode)
                if created:
                    imported += 1
                else:
                    skipped += 1
            except Exception as exc:
                errors.append(f"Row {row.get('_row', '?')}: {exc}")

        if not errors:
            db.session.commit()
        else:
            db.session.rollback()
    except Exception as exc:
        db.session.rollback()
        errors.append(str(exc))

    return {
        "table": spec["label"],
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
    }


def _clear_table(table_key):
    if table_key == "member_ledger":
        MemberLedger.query.delete()
    elif table_key == "journal_lines":
        JournalLine.query.delete()
    elif table_key == "journal_entries":
        JournalLine.query.delete()
        MemberLedger.query.filter(MemberLedger.journal_entry_id.isnot(None)).update(
            {"journal_entry_id": None}, synchronize_session=False
        )
        JournalEntry.query.delete()
    elif table_key == "members":
        MemberLedger.query.delete()
        Member.query.delete()
    elif table_key == "users":
        User.query.filter(User.role != "PlatformAdmin").delete(synchronize_session=False)
    elif table_key == "accounts":
        MemberLedger.query.filter(MemberLedger.journal_entry_id.isnot(None)).update(
            {"journal_entry_id": None}, synchronize_session=False
        )
        JournalLine.query.delete()
        JournalEntry.query.delete()
        Account.query.delete()
    elif table_key == "cooperatives":
        Cooperative.query.delete()
    db.session.flush()


def _cooperative_field_values(row):
    return {
        "name": str(row["name"]).strip(),
        "registration_no": row.get("registration_no"),
        "tin": row.get("tin"),
        "rdo": row.get("rdo"),
        "coop_type": row.get("coop_type"),
        "address": row.get("address"),
        "fiscal_year_end": row.get("fiscal_year_end", "December 31"),
    }


def _sync_registry_from_cooperative_row(row):
    """Keep coop_registry profile in sync with cooperatives import."""
    registry = get_coop_registry()
    if not registry:
        return
    fields = _cooperative_field_values(row)
    registry.name = fields["name"]
    registry.registration_no = fields.get("registration_no")
    registry.tin = fields.get("tin")
    registry.rdo = fields.get("rdo")
    registry.coop_type = fields.get("coop_type")
    registry.address = fields.get("address")
    registry.fiscal_year_end = fields.get("fiscal_year_end") or registry.fiscal_year_end
    if row.get("contact_email"):
        registry.contact_email = str(row.get("contact_email")).strip()


def _import_one_row(table_key, row, mode):
    if table_key == "cooperatives":
        fields = _cooperative_field_values(row)
        existing = Cooperative.query.first()
        if existing and mode == "append":
            for key, value in fields.items():
                if value is not None:
                    setattr(existing, key, value)
            _sync_registry_from_cooperative_row(row)
            return False
        if existing:
            return False
        db.session.add(Cooperative(**fields))
        _sync_registry_from_cooperative_row(row)
        return True

    if table_key == "accounts":
        code = str(row["code"]).strip()
        if Account.query.filter_by(code=code).first():
            return False
        atype = str(row["account_type"]).strip()
        if atype not in ACCOUNT_TYPES:
            raise ValueError(f"invalid account_type '{atype}'")
        nb = str(row["normal_balance"]).strip()
        if nb not in NORMAL_BALANCES:
            raise ValueError("normal_balance must be Debit or Credit")
        db.session.add(Account(
            code=code,
            name=str(row["name"]).strip(),
            account_type=atype,
            category=row.get("category"),
            normal_balance=nb,
            is_active=_parse_bool(row.get("is_active")),
        ))
        return True

    if table_key == "users":
        username = str(row["username"]).strip()
        if User.query.filter_by(username=username).first():
            return False
        role = normalize_role(row.get("role", "Member"))
        if not is_valid_user_role(role):
            raise ValueError(f"invalid role '{row.get('role')}'")
        status = str(row.get("status", "Active")).strip()
        if status not in ("Active", "Inactive"):
            raise ValueError("status must be Active or Inactive")
        db.session.add(User(
            username=username,
            full_name=row.get("full_name"),
            email=row.get("email"),
            role=role,
            status=status,
            password_hash=generate_password_hash(str(row["password"])),
        ))
        return True

    if table_key == "members":
        member_no = str(row["member_no"]).strip()
        if Member.query.filter_by(member_no=member_no).first():
            return False

        row_errors = _validate_member_row(row)
        if row_errors:
            raise ValueError(row_errors[0])

        last_name, first_name, middle_name = _parse_member_names(row)

        paid_up = _parse_decimal(row.get("initial_paid_up_capital") or row.get("share_capital"))
        savings = _parse_decimal(row.get("savings_balance"))
        membership_date = _parse_date(row.get("membership_date"))
        register_entry = _parse_date(row.get("register_entry_date")) or membership_date

        member = Member(
            member_no=member_no,
            last_name=last_name,
            first_name=first_name,
            middle_name=middle_name,
            full_name=build_full_name(last_name, first_name, middle_name),
            tin=row.get("tin"),
            membership_date=membership_date,
            bod_acceptance_resolution=row.get("bod_acceptance_resolution"),
            membership_type=str(row.get("membership_type") or "Regular").strip(),
            initial_shares=_parse_decimal(row.get("initial_shares")),
            initial_subscription_amount=_parse_decimal(row.get("initial_subscription_amount")),
            initial_paid_up_capital=paid_up,
            address=row.get("address"),
            birth_date=_parse_date(row.get("birth_date")),
            gender=row.get("gender"),
            civil_status=row.get("civil_status"),
            highest_education=row.get("highest_education"),
            occupation_income_source=row.get("occupation_income_source"),
            number_of_dependents=_parse_int(row.get("number_of_dependents")),
            religion_social_affiliation=row.get("religion_social_affiliation"),
            annual_income=_parse_decimal(row.get("annual_income")),
            register_entry_date=register_entry,
            termination_date=_parse_date(row.get("termination_date")),
            termination_bod_resolution=row.get("termination_bod_resolution"),
            email=row.get("email"),
            phone=row.get("phone"),
            share_capital=paid_up,
            savings_balance=savings,
            status=str(row.get("status", "Active")).strip(),
        )
        db.session.add(member)
        db.session.flush()
        create_opening_ledger_entries(
            member,
            posted_by="Migration Import",
            share_amount=paid_up,
            savings_amount=savings,
        )
        return True

    if table_key == "journal_entries":
        entry_no = str(row["entry_no"]).strip()
        if JournalEntry.query.filter_by(entry_no=entry_no).first():
            return False
        db.session.add(JournalEntry(
            entry_no=entry_no,
            entry_date=_parse_date(row["entry_date"]),
            description=row.get("description"),
            reference=row.get("reference"),
            posted_by=row.get("posted_by"),
        ))
        return True

    if table_key == "journal_lines":
        entry = JournalEntry.query.filter_by(entry_no=str(row["entry_no"]).strip()).first()
        account = Account.query.filter_by(code=str(row["account_code"]).strip()).first()
        if not entry or not account:
            raise ValueError("entry_no or account_code not found")
        db.session.add(JournalLine(
            entry_id=entry.id,
            account_id=account.id,
            debit=_parse_decimal(row.get("debit")),
            credit=_parse_decimal(row.get("credit")),
            memo=row.get("memo"),
        ))
        return True

    if table_key == "member_ledger":
        member = Member.query.filter_by(member_no=str(row["member_no"]).strip()).first()
        if not member:
            raise ValueError(f"member_no not found: {row.get('member_no')}")
        ltype = str(row["ledger_type"]).strip().lower()
        if ltype not in LEDGER_TYPES:
            raise ValueError("ledger_type must be share, savings, or loan")
        journal_entry_id = None
        if row.get("entry_no"):
            entry = JournalEntry.query.filter_by(entry_no=str(row["entry_no"]).strip()).first()
            if not entry:
                raise ValueError(f"entry_no not found: {row.get('entry_no')}")
            journal_entry_id = entry.id
        db.session.add(MemberLedger(
            member_id=member.id,
            txn_date=_parse_date(row["txn_date"]),
            txn_type=str(row["txn_type"]).strip(),
            ledger_type=ltype,
            reference=row.get("reference"),
            description=row.get("description"),
            debit=_parse_decimal(row.get("debit")),
            credit=_parse_decimal(row.get("credit")),
            posted_by=row.get("posted_by"),
            journal_entry_id=journal_entry_id,
        ))
        return True

    raise ValueError(f"Unknown table: {table_key}")


def import_batch(uploads, mode="append"):
    """Import multiple tables in dependency order."""
    results = []
    for spec in IMPORT_ORDER:
        key = spec["key"]
        if key not in uploads:
            continue
        rows, parse_errors = read_upload_rows(uploads[key], key)
        if parse_errors:
            results.append({
                "table": spec["label"],
                "imported": 0,
                "skipped": 0,
                "errors": parse_errors,
            })
            continue
        if not rows:
            continue
        results.append(import_rows(key, rows, mode=mode))
    return results
